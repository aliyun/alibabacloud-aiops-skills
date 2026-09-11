"""VPC boundary ACL policy backup plugin"""
import json
from .base import AclPolicyPlugin


# Action value mapping
ACL_ACTION_MAP = {
    "accept": "放行",
    "drop": "拦截",
    "log": "观察",
}

# Action value reverse mapping
ACL_ACTION_REVERSE_MAP = {
    "放行": "accept",
    "拦截": "drop",
    "观察": "log",
}


def _format_timestamp(ts):
    """Convert a timestamp to a human-readable format"""
    if not ts:
        return ""
    from datetime import datetime, timezone, timedelta
    try:
        dt = datetime.fromtimestamp(int(ts), tz=timezone(timedelta(hours=8)))
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(ts)


class VpcAclPlugin(AclPolicyPlugin):
    """VPC boundary ACL policy"""

    name = "VPC边界"
    order = 3
    sheet_name = "VPC边界墙ACL"
    api_action = "DescribeVpcFirewallControlPolicy"
    result_key = "Policys"
    group_types = []  # GroupType is not used for ACL policies

    columns = [
        ("策略组ID",        lambda item: item.get("_AclGroupId", "")),
        ("策略组名称",      lambda item: item.get("_AclGroupName", "")),
        ("优先级",          lambda item: item.get("Order", "")),
        ("访问源",          lambda item: item.get("Source", "")),
        ("源地址类型",      lambda item: item.get("SourceType", "")),
        ("源地址簿类型",    lambda item: item.get("SourceGroupType", "")),
        ("目的",            lambda item: item.get("Destination", "")),
        ("目的地址类型",    lambda item: item.get("DestinationType", "")),
        ("目的地址簿类型",  lambda item: item.get("DestinationGroupType", "")),
        ("域名解析类型",    lambda item: item.get("DomainResolveType", "")),
        ("网络传输协议",    lambda item: item.get("Proto", "")),
        ("应用名称列表",    lambda item: ", ".join(item.get("ApplicationNameList", []))),
        ("端口",            lambda item: item.get("DestPort", "")),
        ("端口类型",        lambda item: item.get("DestPortType", "")),
        ("端口地址簿",      lambda item: item.get("DestPortGroup", "")),
        ("动作",            lambda item: ACL_ACTION_MAP.get(item.get("AclAction", ""), item.get("AclAction", ""))),
        ("策略有效期",      lambda item: item.get("RepeatType", "")),
        ("重复周期",        lambda item: ", ".join(str(d) for d in item.get("RepeatDays", [])) if item.get("RepeatDays") else ""),
        ("重复开始时间",    lambda item: item.get("RepeatStartTime", "")),
        ("重复结束时间",    lambda item: item.get("RepeatEndTime", "")),
        ("策略生效开始时间", lambda item: _format_timestamp(item.get("StartTime", 0))),
        ("策略生效结束时间", lambda item: _format_timestamp(item.get("EndTime", 0))),
        ("启用状态",        lambda item: item.get("Release", "")),
        ("描述",            lambda item: item.get("Description", "")),
        ("展开数量",        lambda item: item.get("SpreadCnt", "")),
    ]

    col_widths = [28, 25, 8, 30, 12, 15, 30, 12, 15, 15, 15, 15, 15, 10, 20, 10, 12, 10, 20, 20, 12, 12, 10, 30, 10]

    def custom_fetch(self, ak, sk, endpoint, call_api_fn, page_size=50, security_token=None):
        """
        VPC boundary ACL backup flow:
        1. Query the list of all policy groups (DescribeVpcFirewallAclGroupList)
        2. For each policy group, query ACL policies (VpcFirewallId is set to AclGroupId)
        """
        all_items = []

        # Step 1: get all policy groups
        acl_group_data = call_api_fn(ak, sk, endpoint, "DescribeVpcFirewallAclGroupList", {
            "CurrentPage": "1",
            "PageSize": str(page_size),
            "Lang": "zh",
        }, security_token)

        if not acl_group_data:
            print(f"  [{self.name}] failed to query the policy group list")
            return all_items

        groups = acl_group_data.get("AclGroupList", [])
        print(f"  [{self.name}] found {len(groups)} policy group(s)")

        # Step 2: query ACL policies for each policy group
        for group in groups:
            acl_group_id = group.get("AclGroupId", "")
            acl_group_name = group.get("AclGroupName", "")
            rule_count = group.get("AclRuleCount", 0)

            if not acl_group_id:
                continue

            # Skip policy groups without rules
            if rule_count == 0:
                print(f"    [{self.name}] policy group={acl_group_name} (ID={acl_group_id}), no rules, skipped")
                continue

            page = 1
            while True:
                params = {
                    "CurrentPage": str(page),
                    "PageSize": str(page_size),
                    "VpcFirewallId": acl_group_id,
                    "Lang": "zh",
                }
                data = call_api_fn(ak, sk, endpoint, self.api_action, params, security_token)
                if not data:
                    break

                items = data.get("Policys", [])
                if not items:
                    total = int(data.get("TotalCount", 0))
                    if total > 0:
                        print(f"    [{self.name}] policy group={acl_group_name}, {total} rows in total (no data on this page)")
                    break

                # Attach policy group information to each policy
                for item in items:
                    item["_AclGroupId"] = acl_group_id
                    item["_AclGroupName"] = acl_group_name

                all_items.extend(items)
                total = int(data.get("TotalCount", 0))
                print(f"    [{self.name}] policy group={acl_group_name}, page {page}, {len(items)} rows on this page, accumulated {len(all_items)}/{total}")

                if len(all_items) >= total:
                    break
                page += 1

        return all_items

    # ==================== Restore configuration ====================
    restore_api = "CreateVpcFirewallControlPolicy"
    name_column = "策略组名称"
    uid_column = "AclUuid"

    def restore(self, ak, sk, region, call_api_fn, excel_file, security_token=None,
                vpc_target_groups=None):
        """VPC boundary ACL policy restore (overrides the base class to handle the policy group concept)

        vpc_target_groups: optional, list of target policy group IDs explicitly specified by the
        user (append mode). When specified, it takes priority: regardless of whether the policy
        groups in the backup exist in this account, all policies' resource IDs are rewritten to
        the specified targets and restored in a broadcast manner (covering both policy-group-swap-
        in-same-account and cross-account scenarios). When not specified and the backup policy
        groups are missing in this account, an interactive selection is shown or a hint to use the
        parameter.
        """
        import pandas as pd
        import sys as _sys
        from datetime import datetime, timezone, timedelta

        endpoint = f"cloudfw.{region}.aliyuncs.com"

        # Skip if the sheet does not exist (no sheet is generated for boundaries without data at backup time)
        if self.sheet_name not in pd.ExcelFile(excel_file).sheet_names:
            print(f"  [{self.name}] no sheet '{self.sheet_name}' in the backup file, skipped (no data at backup time)")
            return (0, 0, 0, 0)

        df = pd.read_excel(excel_file, sheet_name=self.sheet_name)
        print(f"  [{self.name}] loaded {len(df)} records")

        if '恢复状态' not in df.columns:
            df['恢复状态'] = '待恢复'
        else:
            df['恢复状态'] = df['恢复状态'].astype(str).replace('nan', '待恢复')

        success_count = 0
        fail_count = 0
        skip_count = 0
        exist_count = 0

        # ===== Target policy group selection (cross-account restore support, multi-select broadcast) =====
        def _available_groups():
            """Query VPC firewall policy groups in this account (DescribeVpcFirewallAclGroupList, paginated)"""
            groups, page = {}, 1
            while True:
                data = call_api_fn(ak, sk, endpoint, "DescribeVpcFirewallAclGroupList", {
                    "CurrentPage": str(page), "PageSize": "50", "Lang": "zh"}, security_token)
                batch = (data or {}).get("AclGroupList", [])
                for g in batch:
                    gid = g.get("AclGroupId")
                    if gid:
                        groups[gid] = g.get("AclGroupName", "") or "(unnamed)"
                total = int((data or {}).get("TotalCount", 0) or 0)
                if not batch or page * 50 >= total:
                    break
                page += 1
            return groups

        source_groups = sorted({str(g) for g in df["策略组ID"].astype(str).unique()
                                if g and g != "nan"})
        available_groups = _available_groups()
        missing_groups = [g for g in source_groups if g not in available_groups]
        remap_groups = None
        if vpc_target_groups:
            # User explicitly specified target policy groups: takes priority, rewrite resource IDs and broadcast-restore (policy-group-swap-in-same-account/cross-account)
            invalid = [g for g in vpc_target_groups if g not in available_groups]
            if invalid:
                print(f"  [{self.name}] [FAIL] specified policy groups do not exist: {', '.join(invalid)}; "
                      f"available policy groups: {', '.join(f'{k}({v})' for k, v in available_groups.items()) if available_groups else '(no available policy group in this account)'}")
                return (0, len(df[df["恢复状态"] != "成功"]), 0, 0)
            remap_groups = list(vpc_target_groups)
            print(f"  [{self.name}] [INFO] rewriting policy group IDs to the specified targets: backup policy groups [{', '.join(source_groups)}] -> target policy groups [{', '.join(remap_groups)}]")
        elif missing_groups:
            print(f"  [{self.name}] [WARN] {len(missing_groups)} policy group(s) in the backup do not exist in this account: {', '.join(missing_groups)}")
            if not available_groups:
                print(f"  [{self.name}] no VPC firewall policy group in this account. Please activate VPC firewall first and retry; this boundary is skipped for now.")
                return (0, len(df[df["恢复状态"] != "成功"]), 0, 0)
            if _sys.stdin.isatty():
                print(f"  Available policy groups in this account (multiple allowed, comma-separated indices):")
                items = list(available_groups.items())
                for i, (gid, gname) in enumerate(items, 1):
                    print(f"    [{i}] {gid} ({gname})")
                while True:
                    sel = input("  Please select the policy group(s) to restore to: ").strip()
                    try:
                        picked = [items[int(t) - 1][0] for t in sel.split(",") if t.strip()]
                        if picked:
                            remap_groups = picked
                            break
                    except (ValueError, IndexError):
                        pass
                    print("  Invalid input, please select again")
            else:
                print(f"  [{self.name}] [FAIL] in a non-interactive environment, use --vpc-target-groups cen-xxx[,cen-yyy] "
                      f"to specify the target policy groups (multiple allowed; policies will be broadcast-restored to each policy group); "
                      f"available policy groups: {', '.join(available_groups)}")
                return (0, len(df[df["恢复状态"] != "成功"]), 0, 0)
            print(f"  [{self.name}] all policies will be restored to target policy groups: {', '.join(remap_groups)}")
        else:
            # No targets specified and all backup policy groups exist in this account: keep the policy group IDs from the backup (same-resource append scenario)
            print(f"  [{self.name}] [INFO] no target policy group specified, using the policy groups from the backup: {', '.join(source_groups)}"
                  f"(to restore to other policy groups, specify them with --vpc-target-groups; "
                  f"to restore the existing configuration of the same policy group, use --mode diff instead)")

        # Iterate in ascending order (NewOrder=-1 inserts at the end, so ascending order keeps the sequence consistent)
        for idx in range(len(df)):
            row = df.iloc[idx]
            record_name = str(row.get("策略组名称", f"record{idx+1}"))
            acl_group_id = str(row.get("策略组ID", ""))
            targets = remap_groups if remap_groups else [acl_group_id]

            # Skip if already restored successfully
            current_status = df.at[df.index[idx], '恢复状态']
            if current_status == '成功':
                skip_count += 1
                continue

            # ===== Build API parameters =====
            api_params = {}

            # VpcFirewallId = policy group ID
            api_params["VpcFirewallId"] = acl_group_id

            # AclAction: accept=accept, log=observe, drop=reject
            action_str = str(row.get("动作", ""))
            acl_action = ACL_ACTION_REVERSE_MAP.get(action_str, action_str)
            api_params["AclAction"] = acl_action

            # Source / SourceType
            api_params["Source"] = str(row.get("访问源", ""))
            api_params["SourceType"] = str(row.get("源地址类型", ""))

            # Destination / DestinationType
            api_params["Destination"] = str(row.get("目的", ""))
            api_params["DestinationType"] = str(row.get("目的地址类型", ""))

            # Proto
            api_params["Proto"] = str(row.get("网络传输协议", ""))

            # ApplicationNameList: flattened format, required (including ANY)
            app_list_str = str(row.get("应用名称列表", ""))
            if app_list_str and app_list_str != "nan":
                apps = [a.strip() for a in app_list_str.split(",") if a.strip()]
                for i, app in enumerate(apps, 1):
                    api_params[f"ApplicationNameList.{i}"] = app

            # DestPortType
            dest_port_type = str(row.get("端口类型", ""))
            api_params["DestPortType"] = dest_port_type

            # DestPort / DestPortGroup
            if dest_port_type == "group":
                dest_port_group = str(row.get("端口地址簿", ""))
                if dest_port_group and dest_port_group != "nan":
                    api_params["DestPortGroup"] = dest_port_group
                api_params["DestPort"] = ""
            else:
                dest_port = str(row.get("端口", ""))
                api_params["DestPort"] = dest_port if dest_port != "nan" else ""

            # NewOrder: insert at the end by default
            api_params["NewOrder"] = "-1"

            # Release: enable status
            release_str = str(row.get("启用状态", ""))
            if release_str == "True" or release_str == "true":
                api_params["Release"] = "true"
            elif release_str == "False" or release_str == "false":
                api_params["Release"] = "false"
            else:
                api_params["Release"] = "true"

            # Description (required)
            description = str(row.get("描述", ""))
            api_params["Description"] = description if description != "nan" else ""

            # RepeatType
            repeat_type = str(row.get("策略有效期", ""))
            if repeat_type and repeat_type != "nan":
                api_params["RepeatType"] = repeat_type

                # Pass time parameters when not Permanent
                if repeat_type != "Permanent":
                    start_time_str = str(row.get("策略生效开始时间", ""))
                    end_time_str = str(row.get("策略生效结束时间", ""))

                    if start_time_str and start_time_str != "nan":
                        try:
                            dt = datetime.strptime(start_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["StartTime"] = str(int(dt.timestamp()))
                        except:
                            pass

                    if end_time_str and end_time_str != "nan":
                        try:
                            dt = datetime.strptime(end_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["EndTime"] = str(int(dt.timestamp()))
                        except:
                            pass

                    repeat_start = str(row.get("重复开始时间", ""))
                    repeat_end = str(row.get("重复结束时间", ""))
                    if repeat_start and repeat_start != "nan":
                        api_params["RepeatStartTime"] = repeat_start
                    if repeat_end and repeat_end != "nan":
                        api_params["RepeatEndTime"] = repeat_end

                    # RepeatDays (Weekly type)
                    repeat_days_str = str(row.get("重复周期", ""))
                    if repeat_days_str and repeat_days_str != "nan":
                        days = [d.strip() for d in repeat_days_str.split(",") if d.strip()]
                        for i, day in enumerate(days, 1):
                            api_params[f"RepeatDays.{i}"] = day
            else:
                # When the policy validity period is empty but effective dates exist, pass RepeatType=None (one-off time range)
                start_time_str = str(row.get("策略生效开始时间", ""))
                end_time_str = str(row.get("策略生效结束时间", ""))
                if (start_time_str and start_time_str != "nan") or (end_time_str and end_time_str != "nan"):
                    api_params["RepeatType"] = "None"
                    if start_time_str and start_time_str != "nan":
                        try:
                            dt = datetime.strptime(start_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["StartTime"] = str(int(dt.timestamp()))
                        except:
                            pass
                    if end_time_str and end_time_str != "nan":
                        try:
                            dt = datetime.strptime(end_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["EndTime"] = str(int(dt.timestamp()))
                        except:
                            pass

            # DomainResolveType: pass it when the destination address type is domain
            dest_type = str(row.get("目的地址类型", ""))
            if dest_type == "domain":
                domain_resolve_type = str(row.get("域名解析类型", ""))
                if domain_resolve_type and domain_resolve_type != "nan":
                    api_params["DomainResolveType"] = domain_resolve_type

            print(f"  [{self.name}] [{idx+1}] restoring: {record_name} (target policy groups: {', '.join(targets)})")
            print(f"    action: {acl_action}, proto: {api_params['Proto']}, source: {api_params['Source']}")

            ok_t, exists_t, fail_t = 0, 0, 0
            fail_msgs = []
            for grp in targets:
                params = dict(api_params)
                params["VpcFirewallId"] = grp
                try:
                    result = call_api_fn(ak, sk, endpoint, self.restore_api, params, security_token)

                    if self.uid_column in result and result[self.uid_column]:
                        print(f"    [OK] {grp} success, {self.uid_column}: {result[self.uid_column]}")
                        ok_t += 1
                    else:
                        code = result.get("Code", "")
                        message = result.get("Message", "")
                        if "already exists" in message.lower() or "Duplicates" in code:
                            print(f"    [WARN] {grp} already exists")
                            exists_t += 1
                        else:
                            print(f"    [FAIL] {grp} failed: {code} - {message}")
                            fail_t += 1
                            fail_msgs.append(f"{grp}:{code}")

                except Exception as e:
                    error_str = str(e)
                    if "already exist" in error_str.lower() or "Duplicates" in error_str:
                        print(f"    [WARN] {grp} already exists")
                        exists_t += 1
                    else:
                        print(f"    [FAIL] {grp} exception: {error_str[:200]}")
                        fail_t += 1
                        fail_msgs.append(f"{grp}:{error_str[:50]}")

            total_t = len(targets)
            if fail_t == 0 and ok_t + exists_t == total_t:
                df.at[df.index[idx], '恢复状态'] = '成功'
                success_count += 1
            elif ok_t + exists_t == 0:
                df.at[df.index[idx], '恢复状态'] = '失败'
                fail_count += 1
            else:
                df.at[df.index[idx], '恢复状态'] = f"部分成功({ok_t+exists_t}/{total_t}，失败:{';'.join(fail_msgs)})"
                fail_count += 1

        # Update the restore status in Excel
        import openpyxl
        wb = openpyxl.load_workbook(excel_file)
        ws = wb[self.sheet_name]
        header_row = [cell.value for cell in ws[1]]
        if '恢复状态' in header_row:
            status_col_idx = header_row.index('恢复状态') + 1
        else:
            status_col_idx = ws.max_column + 1
            ws.cell(row=1, column=status_col_idx, value='恢复状态')

        for i, (_, row) in enumerate(df.iterrows()):
            ws.cell(row=i + 2, column=status_col_idx, value=row.get('恢复状态', '待恢复'))

        wb.save(excel_file)

        print(f"  [{self.name}] restore completed: success {success_count}, failed {fail_count}, skipped {skip_count}, already existing {exist_count}")
        return (success_count, fail_count, skip_count, exist_count)
