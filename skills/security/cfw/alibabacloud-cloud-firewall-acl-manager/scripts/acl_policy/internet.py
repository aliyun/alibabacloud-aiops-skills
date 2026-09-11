"""Internet boundary ACL policy backup plugin"""
import json
from .base import AclPolicyPlugin


# Action value mapping
ACL_ACTION_MAP = {
    "accept": "放行",
    "drop": "拦截",
    "log": "观察",
}

# Action value reverse mapping
ACL_ACTION_REVERSE_MAP = {
    "放行": "accept",
    "拦截": "drop",
    "观察": "log",
}

# Direction value mapping
DIRECTION_MAP = {
    "in": "入方向",
    "out": "出方向",
}

# Direction value reverse mapping
DIRECTION_REVERSE_MAP = {
    "入方向": "in",
    "出方向": "out",
}


def _parse_source_cidrs(item):
    """Parse the CIDR list expanded from the source address book"""
    cidrs = item.get("SourceGroupCidrs", [])
    return "\n".join(cidrs) if cidrs else ""


def _parse_dest_cidrs(item):
    """Parse the CIDR list expanded from the destination address book"""
    cidrs = item.get("DestinationGroupCidrs", [])
    return "\n".join(cidrs) if cidrs else ""


def _parse_dest_ports(item):
    """Parse the port list expanded from the port address book"""
    ports = item.get("DestPortGroupPorts", [])
    return "\n".join(ports) if ports else ""


def _format_timestamp(ts):
    """Convert a timestamp to a human-readable format"""
    if not ts:
        return ""
    from datetime import datetime, timezone, timedelta
    try:
        dt = datetime.fromtimestamp(int(ts), tz=timezone(timedelta(hours=8)))
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(ts)


class InternetAclPlugin(AclPolicyPlugin):
    """Internet boundary ACL policy"""

    name = "互联网边界"
    order = 1
    sheet_name = "互联网边界墙ACL"
    api_action = "DescribeControlPolicy"
    result_key = "Policys"
    group_types = []  # GroupType is not used for ACL policies

    columns = [
        ("方向",            lambda item: DIRECTION_MAP.get(item.get("Direction", ""), item.get("Direction", ""))),
        ("IP版本",          lambda item: item.get("IpVersion", "")),
        ("优先级",          lambda item: item.get("Order", "")),
        ("策略名称",        lambda item: item.get("AclName", "")),
        ("描述",            lambda item: item.get("Description", "")),
        ("源地址",          lambda item: item.get("Source", "")),
        ("源地址类型",      lambda item: item.get("SourceType", "")),
        ("源地址簿类型",    lambda item: item.get("SourceGroupType", "")),
        ("目的地址",        lambda item: item.get("Destination", "")),
        ("目的地址类型",    lambda item: item.get("DestinationType", "")),
        ("目的地址簿类型",  lambda item: item.get("DestinationGroupType", "")),
        ("域名解析类型",    lambda item: item.get("DomainResolveType", "")),
        ("协议",            lambda item: item.get("Proto", "")),
        # ApplicationNameList: when a source policy has no application bound (created by legacy
        # CLI/API, system default, or NDR auto-creation), Describe returns empty; during backup it
        # is normalized to ANY (the server-side standard value for "no specific application bound",
        # verified: after writing ANY, reading it back yields ['ANY'], semantically lossless),
        # ensuring the backup file is complete and self-consistent and reviewable before restore.
        ("应用名称列表",    lambda item: ", ".join(item.get("ApplicationNameList", [])) or "ANY"),
        ("端口",            lambda item: item.get("DestPort", "")),
        ("端口类型",        lambda item: item.get("DestPortType", "")),
        ("端口地址簿",      lambda item: item.get("DestPortGroup", "")),
        ("动作",            lambda item: ACL_ACTION_MAP.get(item.get("AclAction", ""), item.get("AclAction", ""))),
        ("应用模板",        lambda item: item.get("ApplicationTemplate", "")),
        ("Web过滤模板",     lambda item: item.get("WebFilterTemplate", "")),
        ("重复类型",        lambda item: item.get("RepeatType", "")),
        ("重复周期",        lambda item: ", ".join(str(d) for d in item.get("RepeatDays", [])) if item.get("RepeatDays") else ""),
        ("重复开始时间",    lambda item: item.get("RepeatStartTime", "")),
        ("重复结束时间",    lambda item: item.get("RepeatEndTime", "")),
        ("策略生效开始时间", lambda item: _format_timestamp(item.get("StartTime", 0))),
        ("策略生效结束时间", lambda item: _format_timestamp(item.get("EndTime", 0))),
        ("启用状态",        lambda item: item.get("Release", "")),
        ("展开数量",        lambda item: item.get("SpreadCnt", "")),
    ]

    col_widths = [10, 8, 10, 25, 30, 30, 12, 15, 30, 12, 15, 15, 10, 20, 15, 10, 20, 10, 25, 25, 12, 20, 20, 12, 12, 10, 10]

    def extra_params(self):
        return {"Direction": "in"}

    def custom_fetch(self, ak, sk, endpoint, call_api_fn, page_size=50, security_token=None):
        """Query IPv4 and IPv6 policies separately; for each version, query both inbound and outbound directions"""
        all_items = []

        for ip_version in ["4", "6"]:
            ip_label = "IPv4" if ip_version == "4" else "IPv6"
            for direction in ["in", "out"]:
                page = 1
                while True:
                    params = {
                        "CurrentPage": str(page),
                        "PageSize": str(page_size),
                        "Direction": direction,
                        "IpVersion": ip_version,
                        "Lang": "zh",
                    }
                    data = call_api_fn(ak, sk, endpoint, self.api_action, params, security_token)
                    if not data:
                        break

                    items = data.get("Policys", [])
                    if not items:
                        break

                    all_items.extend(items)
                    total = int(data.get("TotalCount", 0))
                    print(f"  [{self.name}] {ip_label}, direction={direction}, page {page}, {len(items)} rows on this page, accumulated {len(all_items)}/{total}")

                    if len(all_items) >= total:
                        break
                    page += 1

        return all_items

    # ==================== Restore configuration ====================
    restore_api = "AddControlPolicy"
    name_column = "策略名称"
    uid_column = "AclUuid"

    def restore(self, ak, sk, region, call_api_fn, excel_file, security_token=None):
        """Internet boundary ACL policy restore (overrides the base class to handle complex parameter logic)"""
        import pandas as pd
        from datetime import datetime, timezone, timedelta

        endpoint = f"cloudfw.{region}.aliyuncs.com"

        # Skip if the sheet does not exist (no sheet is generated for boundaries without data at backup time)
        if self.sheet_name not in pd.ExcelFile(excel_file).sheet_names:
            print(f"  [{self.name}] no sheet '{self.sheet_name}' in the backup file, skipped (no data at backup time)")
            return (0, 0, 0, 0)

        df = pd.read_excel(excel_file, sheet_name=self.sheet_name)
        print(f"  [{self.name}] loaded {len(df)} records")

        if '恢复状态' not in df.columns:
            df['恢复状态'] = '待恢复'
        else:
            df['恢复状态'] = df['恢复状态'].astype(str).replace('nan', '待恢复')

        success_count = 0
        fail_count = 0
        skip_count = 0
        exist_count = 0

        # Iterate in ascending order (NewOrder=-1 inserts at the end, so ascending order keeps the sequence consistent)
        indices = list(range(len(df)))

        for idx in indices:
            row = df.iloc[idx]
            record_name = str(row.get("策略名称", f"record{idx+1}"))

            # Skip if already restored successfully
            current_status = df.at[df.index[idx], '恢复状态']
            if current_status == '成功':
                skip_count += 1
                continue

            # ===== Build API parameters =====
            api_params = {}

            # Direction: inbound=in, outbound=out
            direction_str = str(row.get("方向", ""))
            api_params["Direction"] = DIRECTION_REVERSE_MAP.get(direction_str, direction_str)

            # AclName
            api_params["AclName"] = record_name

            # AclAction: accept=accept, log=observe, drop=reject
            action_str = str(row.get("动作", ""))
            acl_action = ACL_ACTION_REVERSE_MAP.get(action_str, action_str)
            api_params["AclAction"] = acl_action

            # Description
            description = str(row.get("描述", ""))
            api_params["Description"] = description if description != "nan" else ""

            # SourceType / DestinationType
            api_params["SourceType"] = str(row.get("源地址类型", ""))
            api_params["DestinationType"] = str(row.get("目的地址类型", ""))

            # Source / Destination
            api_params["Source"] = str(row.get("源地址", ""))
            api_params["Destination"] = str(row.get("目的地址", ""))

            # Proto
            api_params["Proto"] = str(row.get("协议", ""))

            # ApplicationNameList: flattened format, required (including ANY)
            app_list_str = str(row.get("应用名称列表", ""))
            if app_list_str and app_list_str != "nan":
                apps = [a.strip() for a in app_list_str.split(",") if a.strip()]
                for i, app in enumerate(apps, 1):
                    api_params[f"ApplicationNameList.{i}"] = app
            else:
                # Backward compatibility with legacy backup files (this column may be empty in
                # backups generated before the normalization logic was introduced); new backups
                # normalize it to ANY on the backup side, so this branch is normally not reached
                api_params["ApplicationNameList.1"] = "ANY"

            # DestPortType
            dest_port_type = str(row.get("端口类型", ""))
            api_params["DestPortType"] = dest_port_type

            # DestPort / DestPortGroup
            if dest_port_type == "group":
                # Port address book
                dest_port_group = str(row.get("端口地址簿", ""))
                if dest_port_group and dest_port_group != "nan":
                    api_params["DestPortGroup"] = dest_port_group
                api_params["DestPort"] = ""
            else:
                dest_port = str(row.get("端口", ""))
                api_params["DestPort"] = dest_port if dest_port != "nan" else ""

            # NewOrder: insert at the end by default
            api_params["NewOrder"] = "-1"

            # Release: enable status
            release_str = str(row.get("启用状态", ""))
            if release_str == "True" or release_str == "true":
                api_params["Release"] = "true"
            elif release_str == "False" or release_str == "false":
                api_params["Release"] = "false"
            else:
                api_params["Release"] = "true"

            # IpVersion
            ip_version = str(row.get("IP版本", ""))
            if ip_version and ip_version != "nan":
                api_params["IpVersion"] = ip_version

            # RepeatType
            repeat_type = str(row.get("重复类型", ""))
            if repeat_type and repeat_type != "nan":
                api_params["RepeatType"] = repeat_type

                # Pass time parameters when not Permanent
                if repeat_type != "Permanent":
                    # StartTime / EndTime: convert time strings to timestamps
                    start_time_str = str(row.get("策略生效开始时间", ""))
                    end_time_str = str(row.get("策略生效结束时间", ""))

                    if start_time_str and start_time_str != "nan":
                        try:
                            dt = datetime.strptime(start_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["StartTime"] = str(int(dt.timestamp()))
                        except:
                            pass

                    if end_time_str and end_time_str != "nan":
                        try:
                            dt = datetime.strptime(end_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["EndTime"] = str(int(dt.timestamp()))
                        except:
                            pass

                    # RepeatStartTime / RepeatEndTime
                    repeat_start = str(row.get("重复开始时间", ""))
                    repeat_end = str(row.get("重复结束时间", ""))
                    if repeat_start and repeat_start != "nan":
                        api_params["RepeatStartTime"] = repeat_start
                    if repeat_end and repeat_end != "nan":
                        api_params["RepeatEndTime"] = repeat_end

                    # RepeatDays (Weekly/Monthly types)
                    repeat_days_str = str(row.get("重复周期", ""))
                    if repeat_days_str and repeat_days_str != "nan":
                        days = [d.strip() for d in repeat_days_str.split(",") if d.strip()]
                        for i, day in enumerate(days, 1):
                            api_params[f"RepeatDays.{i}"] = day
            else:
                # When the policy validity period is empty but effective dates exist, pass RepeatType=None (one-off time range)
                start_time_str = str(row.get("策略生效开始时间", ""))
                end_time_str = str(row.get("策略生效结束时间", ""))
                if (start_time_str and start_time_str != "nan") or (end_time_str and end_time_str != "nan"):
                    api_params["RepeatType"] = "None"
                    if start_time_str and start_time_str != "nan":
                        try:
                            dt = datetime.strptime(start_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["StartTime"] = str(int(dt.timestamp()))
                        except:
                            pass
                    if end_time_str and end_time_str != "nan":
                        try:
                            dt = datetime.strptime(end_time_str, "%Y-%m-%d %H:%M:%S")
                            dt = dt.replace(tzinfo=timezone(timedelta(hours=8)))
                            api_params["EndTime"] = str(int(dt.timestamp()))
                        except:
                            pass

            # DomainResolveType: pass it when port type=group and destination address book type=domain
            if dest_port_type == "group":
                dest_group_type = str(row.get("目的地址簿类型", ""))
                if dest_group_type == "domain":
                    domain_resolve_type = str(row.get("域名解析类型", ""))
                    if domain_resolve_type and domain_resolve_type != "nan":
                        api_params["DomainResolveType"] = domain_resolve_type

            # WebFilterTemplate: pass it when AclAction is not drop
            if acl_action != "drop":
                web_filter = str(row.get("Web过滤模板", ""))
                if web_filter and web_filter != "nan":
                    api_params["WebFilterTemplate"] = web_filter

            # ApplicationTemplate: pass it when AclAction is not drop
            if acl_action != "drop":
                app_template = str(row.get("应用模板", ""))
                if app_template and app_template != "nan":
                    api_params["ApplicationTemplate"] = app_template

            print(f"  [{self.name}] [{idx+1}] restoring: {record_name}")
            print(f"    direction: {api_params['Direction']}, action: {acl_action}, proto: {api_params['Proto']}")

            try:
                result = call_api_fn(ak, sk, endpoint, self.restore_api, api_params, security_token)

                if self.uid_column in result and result[self.uid_column]:
                    print(f"    [OK] success, {self.uid_column}: {result[self.uid_column]}")
                    df.at[df.index[idx], '恢复状态'] = '成功'
                    success_count += 1
                else:
                    code = result.get("Code", "")
                    message = result.get("Message", "")
                    if "ErrorAddressGroupExist" in code or "already exists" in message.lower() or "Duplicates" in code:
                        print(f"    [WARN] already exists")
                        df.at[df.index[idx], '恢复状态'] = '已存在'
                        exist_count += 1
                    else:
                        print(f"    [FAIL] failed: {code} - {message}")
                        df.at[df.index[idx], '恢复状态'] = '失败'
                        fail_count += 1

            except Exception as e:
                error_str = str(e)
                if "ErrorAddressGroupExist" in error_str or "already exist" in error_str.lower() or "Duplicates" in error_str:
                    print(f"    [WARN] already exists")
                    df.at[df.index[idx], '恢复状态'] = '已存在'
                    exist_count += 1
                else:
                    print(f"    [FAIL] exception: {error_str[:200]}")
                    df.at[df.index[idx], '恢复状态'] = '失败'
                    fail_count += 1

        # Update the restore status in Excel
        import openpyxl
        wb = openpyxl.load_workbook(excel_file)
        ws = wb[self.sheet_name]
        header_row = [cell.value for cell in ws[1]]
        if '恢复状态' in header_row:
            status_col_idx = header_row.index('恢复状态') + 1
        else:
            status_col_idx = ws.max_column + 1
            ws.cell(row=1, column=status_col_idx, value='恢复状态')

        for i, (_, row) in enumerate(df.iterrows()):
            ws.cell(row=i + 2, column=status_col_idx, value=row.get('恢复状态', '待恢复'))

        wb.save(excel_file)

        print(f"  [{self.name}] restore completed: success {success_count}, failed {fail_count}, skipped {skip_count}, already existing {exist_count}")
        return (success_count, fail_count, skip_count, exist_count)
