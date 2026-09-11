"""Core engine - API calls + Excel generation + plugin dispatch (generic, supports both address books and ACL policies)"""
import hashlib
import hmac
import base64
import os
import sys
import uuid
from datetime import datetime, timezone
from urllib.parse import quote, urlencode
from pathlib import Path

import requests
from alibabacloud_credentials.client import Client as CredentialClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

PAGE_SIZE = 50
SKILL_NAME = "alibabacloud-cloud-firewall-acl-manager"

# ==================== Credentials (default credential chain, no explicit AK/SK handling) ====================

def load_default_credentials():
    """Obtain credentials via the default credential chain of the official Alibaba Cloud Credentials SDK.

    The script does not parse any credential file or read explicit AK/SK:
    credential resolution is fully delegated to the default provider chain
    of alibabacloud_credentials (env vars -> credentials file -> aliyun CLI config, chosen by the SDK).
    """
    try:
        cred = CredentialClient().get_credential()
    except Exception as e:
        print(f"Failed to initialize Alibaba Cloud default credential chain: {e}")
        sys.exit(1)
    ak = getattr(cred, "access_key_id", None) if cred else None
    sk = getattr(cred, "access_key_secret", None) if cred else None
    token = getattr(cred, "security_token", None) if cred else None
    if not ak or not sk:
        print("Default credential chain returned no valid credentials, please configure Alibaba Cloud"
              " credentials first (aliyun configure or credentials file)")
        sys.exit(1)
    return ak, sk, (token or None)


# ==================== Observability (UA injection, once per session) ====================

_SESSION_ID = None

def get_session_id():
    """Generate/reuse a session identifier: 32-char lowercase hex, generated once per session and reused throughout"""
    global _SESSION_ID
    if _SESSION_ID is None:
        _SESSION_ID = uuid.uuid4().hex
    return _SESSION_ID

def get_user_agent():
    """UA template: AlibabaCloud-Agent-Skills/<skill-name>/{session-id}"""
    env_ua = os.environ.get("ALIBABA_CLOUD_USER_AGENT")
    if env_ua:
        return env_ua
    return f"AlibabaCloud-Agent-Skills/{SKILL_NAME}/{get_session_id()}"



def sign_request(params, access_key_secret, method="GET"):
    """Alibaba Cloud API V1 signature"""
    sorted_params = sorted(params.items())
    query_string = urlencode([(k, v) for k, v in sorted_params], quote_via=quote)
    string_to_sign = f"{method}&{quote('/', safe='')}&{quote(query_string, safe='')}"
    signing_key = access_key_secret + "&"
    return base64.b64encode(
        hmac.new(signing_key.encode('utf-8'), string_to_sign.encode('utf-8'), hashlib.sha1).digest()
    ).decode('utf-8')


def call_api(ak, sk, endpoint, action, extra_params, security_token=None, method="GET"):
    """Call Alibaba Cloud API (V1 signature, supports GET/POST)"""
    params = {
        "Action": action,
        "Version": "2017-12-07",
        "Format": "JSON",
        "AccessKeyId": ak,
        "SignatureMethod": "HMAC-SHA1",
        "Timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "SignatureVersion": "1.0",
        "SignatureNonce": str(uuid.uuid4()),
    }
    if security_token:
        params["SecurityToken"] = security_token
    params.update(extra_params)
    params["Signature"] = sign_request(params, sk, method)

    if method == "POST":
        resp = requests.post(f"https://{endpoint}/", data=params, timeout=30,
                             headers={"User-Agent": get_user_agent()})
    else:
        resp = requests.get(f"https://{endpoint}/", params=params, timeout=30,
                            headers={"User-Agent": get_user_agent()})
    
    # Try to parse the response body and return detailed info even on error
    try:
        result = resp.json()
        if resp.status_code >= 400:
            # Return error info instead of raising an exception
            return result
        return result
    except:
        # Raise HTTP error if JSON cannot be parsed
        resp.raise_for_status()


def write_sheet(workbook, plugin, items):
    """Write data of one plugin into an Excel sheet (generic, compatible with both address book and ACL policy plugins)"""
    ws = workbook.create_sheet(title=plugin.sheet_name[:31])

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"),
                         top=Side(style="thin"), bottom=Side(style="thin"))

    # Write header
    for ci, (col_name, _) in enumerate(plugin.columns, 1):
        cell = ws.cell(row=1, column=ci, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Write data
    for ri, item in enumerate(items, 2):
        for ci, (col_name, getter) in enumerate(plugin.columns, 1):
            value = getter(item)
            # For columns containing UID or account, convert large numbers to string to avoid scientific notation
            if ('UID' in col_name or '账号' in col_name) and isinstance(value, (int, float)) and abs(value) >= 1e11:
                value = str(int(value))
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border

    # Column widths
    if plugin.col_widths:
        for ci, width in enumerate(plugin.col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = width

    # Freeze first row
    ws.freeze_panes = "A2"


class BackupEngine:
    """Backup core engine (generic version)"""

    def __init__(self, ak, sk, region, security_token=None):
        self.ak = ak
        self.sk = sk
        self.endpoint = f"cloudfw.{region}.aliyuncs.com"
        self.security_token = security_token

    def fetch(self, plugin):
        """Fetch all data of one plugin (generic, compatible with both address book and ACL policy plugins)"""
        
        # Check whether the plugin has custom fetch logic first
        custom_items = plugin.custom_fetch(
            self.ak, self.sk, self.endpoint, 
            call_api, PAGE_SIZE, self.security_token
        )
        if custom_items is not None:
            # Use custom-fetched data and skip the default logic below
            filtered = plugin.post_process(custom_items)
            skipped = len(custom_items) - len(filtered)
            if skipped:
                print(f"  [{plugin.name}] filtered {skipped} records")
            return filtered
        
        all_items = []
        page = 1

        # Build query params: group_types optional (address books have it, ACL does not)
        extra_base = {}
        group_types = getattr(plugin, 'group_types', [])
        if group_types:
            if len(group_types) == 1:
                extra_base["GroupType"] = group_types[0]
            else:
                for i, gt in enumerate(group_types, 1):
                    extra_base[f"GroupTypes.{i}"] = gt
        extra_base.update(plugin.extra_params())

        while True:
            data = call_api(self.ak, self.sk, self.endpoint,
                           plugin.api_action, {
                               plugin.page_no_key: str(page),
                               "PageSize": str(PAGE_SIZE),
                               "Lang": "zh",
                               **extra_base,
                           }, self.security_token)
            total = int(data.get("TotalCount", 0))
            items = data.get(plugin.result_key, [])
            all_items.extend(items)
            print(f"  [{plugin.name}] page {page}, {len(items)} records this page, cumulative {len(all_items)}/{total}")
            if len(all_items) >= total:
                break
            page += 1

        # Plugin post-processing
        filtered = plugin.post_process(all_items)
        skipped = len(all_items) - len(filtered)
        if skipped:
            print(f"  [{plugin.name}] filtered {skipped} records")
        return filtered

    def backup_to_excel(self, plugins, output, append=False):
        """Execute backup: fetch one by one -> one sheet per type -> save Excel
        
        Args:
            plugins: list of plugin objects
            output: output file path
            append: whether to append to an existing file (True loads the existing file and appends new sheets)
        """
        if not plugins:
            print("[FAIL] No matching plugins found")
            return

        # Append mode: load existing file
        if append and Path(output).exists():
            wb = load_workbook(output)
        else:
            wb = Workbook()
            wb.remove(wb.active)  # Remove default empty sheet

        total_count = 0

        for plugin in plugins:
            print(f"\n[{plugin.name}] fetching...")
            items = self.fetch(plugin)
            if items:
                # If a sheet with the same name already exists, delete it first (avoid openpyxl auto-appending numeric suffix)
                if plugin.sheet_name in wb.sheetnames:
                    del wb[plugin.sheet_name]
                write_sheet(wb, plugin, items)
                total_count += len(items)
                print(f"  [OK] fetched {len(items)} records")
            else:
                print(f"  [INFO] no data")

        if wb.sheetnames:
            wb.save(output)
            print(f"\n[OK] wrote {total_count} records in total, saved to: {output}")
        else:
            print("\n[WARN] no data to back up")

    def run(self, plugin_names, output):
        """Backward compatibility: accept a list of plugin names and call backup_to_excel"""
        from address_book import get_plugins
        plugins = get_plugins(plugin_names)
        self.backup_to_excel(plugins, output)
